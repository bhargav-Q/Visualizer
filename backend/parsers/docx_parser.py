from fastapi import UploadFile
import io
import zipfile
import re
import xml.etree.ElementTree as ET
import logging

logger = logging.getLogger(__name__)

# Lazy load RapidOCR engine singleton
_ocr_engine = None

def get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as e:
            logger.warning(f"Could not load RapidOCR engine: {e}")
            _ocr_engine = False
    return _ocr_engine if _ocr_engine is not False else None

def parse_docx(file: UploadFile) -> dict:
    """
    Reads an uploaded .docx file and extracts text and metadata.
    Uses python-docx for document structure (paragraphs & tables) with automatic
    in-memory XML ampersand sanitization and image OCR scanning.
    """
    file.file.seek(0)
    contents = file.file.read()
    file.file.seek(0)

    
    full_text = []
    ocr_image_texts = []
    structured_tsv_parts = []
    
    # Sanitize XML in zip archive to prevent unescaped ampersand parse errors in docx
    sanitized_buf = io.BytesIO()
    try:
        with zipfile.ZipFile(io.BytesIO(contents), 'r') as in_zip, zipfile.ZipFile(sanitized_buf, 'w') as out_zip:
            for item in in_zip.infolist():
                data = in_zip.read(item.filename)
                if item.filename == 'word/document.xml':
                    raw_xml = data.decode('utf-8', errors='ignore')
                    fixed_xml = re.sub(r'&(?![a-zA-Z0-9#]+;)', '&amp;', raw_xml)
                    data = fixed_xml.encode('utf-8')
                out_zip.writestr(item, data)
        sanitized_buf.seek(0)
    except Exception:
        sanitized_buf = io.BytesIO(contents)

    # Method 1: Try python-docx on sanitized buffer
    try:
        import docx
        doc = docx.Document(sanitized_buf)
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                full_text.append(p.text.strip())
        
        # Extract tables as formatted text lines
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if row_cells:
                    full_text.append(" | ".join(row_cells))
    except Exception as docx_err:
        logger.warning(f"python-docx parsing failed, falling back to XML extraction: {docx_err}")
        full_text = []


    # Method 2: Robust XML parsing with unescaped ampersand sanitization
    if not full_text:
        try:
            with zipfile.ZipFile(io.BytesIO(contents)) as zf:
                if 'word/document.xml' in zf.namelist():
                    raw_xml = zf.read('word/document.xml').decode('utf-8', errors='ignore')
                    # Sanitize unescaped ampersands (e.g. "Sales & Marketing" -> "Sales &amp; Marketing")
                    sanitized_xml = re.sub(r'&(?![a-zA-Z0-9#]+;)', '&amp;', raw_xml)
                    tree = ET.fromstring(sanitized_xml)
                    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    
                    for paragraph in tree.findall('.//w:p', ns):
                        texts = [node.text for node in paragraph.findall('.//w:t', ns) if node.text]
                        if texts:
                            full_text.append("".join(texts))
        except Exception as xml_err:
            logger.warning(f"DOCX XML parse error: {xml_err}")


    # Inspect embedded images in word/media/ for table screenshots / OCR text
    try:
        with zipfile.ZipFile(io.BytesIO(contents)) as zf:
            media_files = [f for f in zf.namelist() if f.startswith('word/media/')]
            ocr_engine = get_ocr_engine()
            
            if media_files and ocr_engine:
                for media_name in media_files:
                    ext = media_name.split('.')[-1].lower()
                    if ext in ['png', 'jpg', 'jpeg', 'webp', 'tiff']:
                        try:
                            img_bytes = zf.read(media_name)
                            result, _ = ocr_engine(img_bytes)
                            if result:
                                # Problem 4 Fix: Use spatial grid reconstruction instead of flat word join.
                                # This preserves table column alignment from embedded screenshot images.
                                try:
                                    from parsers.spatial_grid import reconstruct_grid_from_ocr
                                    grid_rows, tsv_output = reconstruct_grid_from_ocr(result)
                                    if grid_rows and tsv_output.strip():
                                        structured_tsv_parts.append(tsv_output.strip())
                                        ocr_image_texts.append(f"--- Embedded Image ({media_name}) ---\n" + tsv_output.strip())
                                    else:
                                        img_text = " ".join([res[1] for res in result]).strip()
                                        if img_text and img_text not in "\n".join(full_text):
                                            ocr_image_texts.append(f"--- Embedded Image ({media_name}) ---\n" + img_text)
                                except Exception:
                                    img_text = " ".join([res[1] for res in result]).strip()
                                    if img_text and img_text not in "\n".join(full_text):
                                        ocr_image_texts.append(f"--- Embedded Image ({media_name}) ---\n" + img_text)
                        except Exception as ocr_err:
                            logger.warning(f"OCR error on DOCX media {media_name}: {ocr_err}")
    except Exception as zip_err:
        logger.warning(f"Zip extraction for DOCX media failed: {zip_err}")

    all_combined = full_text + ocr_image_texts
    return {
        "text": "\n\n".join(all_combined),
        "structured_tsv": "\n".join(structured_tsv_parts),
        "paragraph_count": len(full_text)
    }


