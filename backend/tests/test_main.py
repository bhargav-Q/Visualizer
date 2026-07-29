import pytest
from io import BytesIO

@pytest.fixture(autouse=True)
def mock_db_operations(mocker):
    """Automatically mock DuckDB calls in test_main.py to eliminate disk file locking during tests."""
    mocker.patch("engine.db.save_cached_analytics", return_value=None)
    mocker.patch("engine.db.get_cached_analytics", return_value=None)
    mocker.patch("engine.db.save_document_analytics", return_value=0)

def test_upload_invalid_file_type(client):
    """Test that an unsupported file type (.exe) is rejected."""
    file_content = b"Some binary executable data"
    files = {"file": ("test.exe", file_content, "application/octet-stream")}
    response = client.post("/api/upload", files=files)
    
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]

def test_upload_file_too_large(client, mocker):
    """Test that a >16MB file is rejected."""
    large_content = b"0" * (16 * 1024 * 1024 + 1)
    files = {"file": ("large.pdf", large_content, "application/pdf")}
    response = client.post("/api/upload", files=files)
    
    assert response.status_code == 413
    assert "File too large" in response.json()["detail"]

def test_upload_valid_pdf_endpoint(client, mocker):
    """Test uploading a valid PDF."""
    import io
    try:
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page(width=612, height=792)
        page.insert_text((72, 72), "Mocked PDF text and content")
        buf = io.BytesIO()
        doc.save(buf)
        doc.close()
        file_content = buf.getvalue()
    except Exception:
        file_content = b"%PDF-1.4 Mock PDF Stream %%EOF"

    files = {"file": ("test.pdf", file_content, "application/pdf")}
    response = client.post("/api/upload", files=files)
    
    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "pdf"
    assert data["data_category"] in ("text", "qualitative_document", "mixed")
    assert data["text"] is not None


def test_upload_valid_csv_endpoint(client, mocker):
    """Test uploading a valid CSV."""
    file_content = b"Date,Val\n2025-01-01,10.0"
    files = {"file": ("test.csv", file_content, "text/csv")}
    response = client.post("/api/upload", files=files)

    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "csv"
    assert data["data_category"] == "tabular"
    assert data["tabular"]["row_count"] == 1


def test_upload_valid_xlsx_endpoint(client, mocker):
    """Test uploading a valid XLSX spreadsheet."""
    import io
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Region", "Sales"])
        ws.append(["North", 500.0])
        buf = io.BytesIO()
        wb.save(buf)
        file_content = buf.getvalue()
    except Exception:
        pytest.skip("openpyxl missing")

    files = {"file": ("sales.xlsx", file_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = client.post("/api/upload", files=files)

    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "xlsx"
    assert data["data_category"] == "tabular"
    assert data["tabular"]["row_count"] == 1


def test_upload_valid_docx_endpoint(client, mocker):
    """Test uploading a valid DOCX file."""
    import io
    try:
        import docx
        doc = docx.Document()
        doc.add_paragraph("Docx sample text")
        buf = io.BytesIO()
        doc.save(buf)
        file_content = buf.getvalue()
    except Exception:
        pytest.skip("python-docx missing")

    files = {"file": ("report.docx", file_content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/upload", files=files)

    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "docx"
    assert data["data_category"] in ("text", "qualitative_document", "mixed")


def test_upload_valid_txt_endpoint(client, mocker):
    """Test uploading a valid TXT document."""
    file_content = b"Simple text content for document profiling."
    files = {"file": ("notes.txt", file_content, "text/plain")}
    response = client.post("/api/upload", files=files)

    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "txt"
    assert data["data_category"] in ("text", "qualitative_document")
    assert data["text"]["summary"] is not None

def test_get_document_metrics_not_found(client, mocker):
    """Test retrieving metrics for a file that does not exist in the database."""
    mocker.patch("engine.db.get_metrics_by_file", return_value=[])
    response = client.get("/api/documents/non_existent.pdf/metrics")
    assert response.status_code == 404
    assert "No metrics found" in response.json()["detail"]

def test_get_document_metrics_success(client, mocker):
    """Test retrieving metrics successfully from the DuckDB database."""
    mock_metrics = [
        {
            "id": "1",
            "file_name": "test.pdf",
            "data_type": "metric",
            "category": "Revenue",
            "metric_value": 1000.0,
            "unit": "USD",
            "context_snippet": "Revenue was 1000 USD",
            "page_number": 1,
            "bbox": [10.0, 20.0, 30.0, 40.0],
            "page_width": 612.0,
            "page_height": 792.0,
            "created_at": "2026-07-23 12:00:00"
        }
    ]
    mocker.patch("engine.db.get_metrics_by_file", return_value=mock_metrics)
    response = client.get("/api/documents/test.pdf/metrics")
    assert response.status_code == 200
    data = response.json()
    assert len(data) == 1
    assert data[0]["category"] == "Revenue"
    assert data[0]["page_width"] == 612.0

def test_get_document_pdf_not_found(client, mocker):
    """Test retrieving PDF that does not exist on disk."""
    from pathlib import Path
    mocker.patch("main.DOCUMENTS_DIR", new=Path("/dummy_path"))
    response = client.get("/api/documents/non_existent.pdf/pdf")
    assert response.status_code == 404
    assert "Source PDF file not found" in response.json()["detail"]

